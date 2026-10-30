#!/usr/bin/env python3
"""
skillcrawl.py - engine for the /skillcrawl Claude Code plugin.

This script does the mechanical work (crawling EVERY page, batching 5 pages at a
time, reading the source document - Word / PDF / image / scanned PDF, state
tracking, developer fix reports, re-check diffing). Claude does the judgment work
(spelling / grammar / punctuation / content-vs-document) by reading the batch
files this script produces and writing an issues.json for each batch.

Run from your project folder (where skillcrawl.config.json lives).

Commands
  setup --url U --doc D     create / update skillcrawl.config.json
  auto [--fresh]            decide what to do next (first run / resume / re-check)
  next-batch                fetch the next 5 pending pages into .skillcrawl/batches/...
  submit <batch_dir>        read <batch_dir>/issues.json, update state, write batch csv
  doc-ready                 confirm image/scanned document pages were transcribed
  pages                     list every page checked in this run (name + url)
  coverage-submit           read .skillcrawl/coverage.json (document page -> website page)
  finalize                  close the run, write FIX_REPORT.html / .md / .csv
  status                    show progress
"""
import argparse
import difflib
import csv
import hashlib
import html as htmllib
import json
import os
import re
import shutil
import subprocess
import sys
import time
import xml.etree.ElementTree as ET
from collections import deque
from datetime import datetime, timezone
from urllib.parse import urljoin, urlparse, urldefrag, unquote

import asyncio  # noqa: E402
import threading  # noqa: E402
HERE = os.path.dirname(os.path.abspath(__file__))


def _ensure_deps():
    """Install the Python packages on first use so non-technical users never have to."""
    try:
        import requests  # noqa: F401
        import bs4  # noqa: F401
        return
    except ImportError:
        pass
    if os.environ.get("SKILLCRAWL_BOOT"):
        sys.exit("Could not install required packages. Run:  pip install -r "
                 + os.path.join(HERE, "..", "requirements.txt"))
    print("First-time setup: installing required Python packages (one minute)...", flush=True)
    req = os.path.join(HERE, "..", "requirements.txt")
    ok = False
    for extra in ([], ["--user"]):
        if subprocess.call([sys.executable, "-m", "pip", "install", "--quiet"] + extra
                           + ["-r", req]) == 0:
            ok = True
            break
    if not ok:  # one broken optional package must not block the rest
        for line in open(req, encoding="utf-8"):
            pkg = line.split("#")[0].strip()
            if pkg:
                subprocess.call([sys.executable, "-m", "pip", "install", "--quiet", "--user", pkg])
    os.environ["SKILLCRAWL_BOOT"] = "1"
    os.execv(sys.executable, [sys.executable] + sys.argv)


_ensure_deps()
import requests  # noqa: E402
from bs4 import BeautifulSoup, NavigableString  # noqa: E402

sys.path.insert(0, HERE)
import report  # noqa: E402
from report import SEV_ORDER, issue_row  # noqa: E402

CONFIG_NAME = "skillcrawl.config.json"
WORK = ".skillcrawl"
STATE_PATH = os.path.join(WORK, "state.json")
DOC_TXT = os.path.join(WORK, "doc.txt")
IMG_DIR = os.path.join(WORK, "doc_images")
COVERAGE_PATH = os.path.join(WORK, "coverage.json")
UA = "SkillCrawl/0.2 (content-checker)"
SKIP_EXT = (".jpg", ".jpeg", ".png", ".gif", ".svg", ".webp", ".ico", ".css", ".js",
            ".json", ".xml", ".pdf", ".zip", ".gz", ".mp4", ".mp3", ".woff", ".woff2",
            ".ttf", ".eot", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".rss")
IMG_EXT = (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff")
SUPPORTED_DOC_EXT = (".pdf", ".docx", ".doc", ".xlsx", ".xlsm", ".xls", ".csv", ".txt", ".md",
                     ".html", ".htm") + IMG_EXT
ISSUE_TYPES = {"spelling", "grammar", "punctuation", "content_mismatch", "missing_content",
               "missing_page", "consistency", "broken_page", "other"}
TYPE_LABEL = {"spelling": "Spelling", "grammar": "Grammar", "punctuation": "Punctuation",
              "content_mismatch": "Does not match document", "missing_content": "Content missing",
              "missing_page": "Page missing from website", "consistency": "Inconsistent",
              "broken_page": "Page broken / not loading", "other": "Other"}
STATUS_LABEL = {"new": "NEW", "still_open": "STILL OPEN", "regressed": "REOPENED"}
MISSING_PAGE_URL = "(not found on website)"
TRANSCRIBE_RE = re.compile(r"\[\[TRANSCRIBE: ([^\]]+)\]\]")

DEFAULT_CONFIG = {
    "document": "",
    "url": "",
    "batch_size": 5,
    "max_pages": 300,
    "max_depth": 6,
    "use_sitemap": True,
    "respect_robots": True,
    "render_js": False,
    "expand_collapsed": True,
    "strip_query": True,
    "delay_seconds": 0.2,
    "concurrency": 4,
    "timeout_seconds": 20,
    "page_urls": [],
    "check_all_pages": False,
    "match_depth": 1,
    "match_pages": 60,
    "exclude_patterns": ["/wp-admin", "/wp-login", "/login", "/cart", "/checkout"],
    "ignore_terms": [],
    "report_dir": "./skillcrawl-reports",
    "notify_webhook": "",
}


# ----------------------------------------------------------------- helpers
def now():
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def norm(s):
    s = (s or "").replace("’", "'").replace("‘", "'")
    s = s.replace("“", '"').replace("”", '"')
    return re.sub(r"\s+", " ", s).strip()


def sha(s):
    return hashlib.sha256(s.encode("utf-8", "replace")).hexdigest()


def die(msg):
    print("ERROR: " + msg)
    sys.exit(1)


def load_config(overrides=None):
    if not os.path.exists(CONFIG_NAME):
        die(f"{CONFIG_NAME} not found in {os.getcwd()}. Run /skillcrawl and give the "
            f"website link and document.")
    with open(CONFIG_NAME, encoding="utf-8") as f:
        user = json.load(f)
    cfg = dict(DEFAULT_CONFIG)
    cfg.update(user)
    for k, v in (overrides or {}).items():
        if v:
            cfg[k] = v
    return cfg


def load_state():
    if not os.path.exists(STATE_PATH):
        die("No previous run found. Run /skillcrawl first.")
    with open(STATE_PATH, encoding="utf-8") as f:
        return json.load(f)


def save_state(state):
    os.makedirs(WORK, exist_ok=True)
    tmp = STATE_PATH + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(state, f, indent=2, ensure_ascii=False)
    os.replace(tmp, STATE_PATH)


# ------------------------------------------------------------- document
def _png_for_transcription(img, n):
    """Save a PIL image as a PNG (max 2200px) Claude can read; return its path."""
    os.makedirs(IMG_DIR, exist_ok=True)
    img = img.convert("RGB")
    img.thumbnail((2200, 2200))
    p = os.path.join(IMG_DIR, f"{n:03d}.png")
    img.save(p, "PNG")
    return p


def _extract_pdf(path, img_counter):
    try:
        from pypdf import PdfReader
    except ImportError:
        die("pypdf missing. Run: pip install pypdf")
    reader = PdfReader(path)
    parts, doc = [], None
    for i, pg in enumerate(reader.pages):
        txt = (pg.extract_text() or "").strip()
        if len(txt) >= 20:
            parts.append(f"[Page {i + 1}]\n{txt}")
            continue
        # scanned / image-only page -> render it so Claude can read it
        try:
            import fitz  # PyMuPDF
            from PIL import Image
        except ImportError:
            die("This PDF has scanned (image-only) pages. Run: pip install pymupdf pillow")
        if doc is None:
            doc = fitz.open(path)
        pix = doc[i].get_pixmap(matrix=fitz.Matrix(1.6, 1.6))
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        img_counter[0] += 1
        parts.append(f"[Page {i + 1}]\n[[TRANSCRIBE: {_png_for_transcription(img, img_counter[0])}]]")
    return "\n\n".join(parts)


def _extract_docx(path):
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        die("python-docx missing. Run: pip install python-docx")
    d = docx.Document(path)
    parts = []
    for child in d.element.body.iterchildren():  # keeps paragraphs and tables in order
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, d)
            t = p.text
            if not t.strip():
                parts.append("")
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.lower().startswith("heading") or style.lower() == "title":
                t = "# " + t
            parts.append(t)
        elif tag == "tbl":
            for row in Table(child, d).rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _extract_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        die("openpyxl missing. Run: pip install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# [Sheet: {ws.title}]")
        for row in ws.iter_rows():
            cells = []
            for c in row:
                v = "" if c.value is None else str(c.value).strip()
                link = c.hyperlink.target if c.hyperlink and c.hyperlink.target else ""
                cells.append(f"{v} <{link}>" if link and link not in v else v)
            if any(cells):
                parts.append(" | ".join(cells).rstrip(" |"))
        parts.append("")
    return "\n".join(parts)


def _extract_csv(path):
    parts = []
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
        for row in csv.reader(f):
            cells = [c.strip() for c in row]
            if any(cells):
                parts.append(" | ".join(cells).rstrip(" |"))
    return "\n".join(parts)


def _extract_one(path, img_counter):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _extract_xlsx(path)
    if ext == ".xls":
        die("Old .xls files are not supported. Open it in Excel and Save As .xlsx (or CSV).")
    if ext == ".csv":
        return _extract_csv(path)
    if ext == ".pdf":
        return _extract_pdf(path, img_counter)
    if ext == ".docx":
        return _extract_docx(path)
    if ext == ".doc":
        die("Old .doc files are not supported. Open it in Word and Save As .docx (or PDF).")
    if ext in IMG_EXT:
        try:
            from PIL import Image
        except ImportError:
            die("Pillow missing. Run: pip install pillow")
        img_counter[0] += 1
        png = _png_for_transcription(Image.open(path), img_counter[0])
        return f"[Page {img_counter[0]}]\n[[TRANSCRIBE: {png}]]"
    if ext in (".html", ".htm"):
        with open(path, encoding="utf-8", errors="replace") as f:
            return BeautifulSoup(f.read(), "html.parser").get_text("\n", strip=True)
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


_DOC_CACHE = {}


def _google_export_url(u):
    """Turn a Google Docs / Sheets / Slides / Drive share link into a download link."""
    m = re.search(r"docs\.google\.com/(spreadsheets|document|presentation)/d/([\w-]+)", u)
    if m:
        kind, fid = m.groups()
        fmt = {"spreadsheets": "xlsx", "document": "docx", "presentation": "pdf"}[kind]
        return f"https://docs.google.com/{kind}/d/{fid}/export?format={fmt}", "." + fmt
    m = re.search(r"drive\.google\.com/(?:file/d/|open\?id=|uc\?(?:export=download&)?id=)([\w-]+)", u)
    if m:
        return f"https://drive.google.com/uc?export=download&id={m.group(1)}", ""
    return u, ""


def doc_sources(cfg):
    d = cfg.get("document")
    return [os.path.expanduser(x) for x in (d if isinstance(d, list) else [d]) if x]


def doc_display(cfg):
    """Short, human-friendly list of the shared documents (file names, not full paths)."""
    names = []
    for x in doc_sources(cfg):
        if re.match(r"^https?://", x, re.I):
            names.append(x if len(x) < 60 else x[:57] + "...")
        else:
            names.append(os.path.basename(x.rstrip("/\\")) or x)
    return ", ".join(names)


def safe_label(label):
    return re.sub(r"[^\w.-]+", "_", label).strip("_") or "document"


_DL_LOCK = threading.Lock()


def _download_link(src):
    """Download a document link (Google Sheet / Doc / Drive file / direct link) fresh each run."""
    if src in _DOC_CACHE:
        return _DOC_CACHE[src]
    url, ext = _google_export_url(src)
    try:
        r = requests.get(url, headers={"User-Agent": UA}, timeout=60, allow_redirects=True)
    except Exception as e:
        die(f"Could not download the document link ({type(e).__name__}). Check the link and "
            f"your internet connection.")
    ctype = r.headers.get("content-type", "").lower()
    if r.status_code != 200 or ("text/html" in ctype and "accounts.google" in r.url + r.text[:3000]) \
            or (ext and "text/html" in ctype):
        # private link: use the text Claude in Chrome copied from the signed-in browser, if present
        cached = os.path.join(WORK, "browser_docs", "doc_" + sha(src)[:10] + ".txt")
        if os.path.exists(cached):
            _DOC_CACHE[src] = cached
            return cached
        print(f"DOC_PRIVATE={src}")
        print(f"BROWSER_DOC_FILE={cached}")
        die("The document link needs permission. In Google, click Share -> General access -> "
            "'Anyone with the link' (Viewer) and try again - or download the file "
            "(File -> Download) and give me the file instead.")
    cd = r.headers.get("content-disposition", "")
    fname = ""
    m = re.search(r"filename\*=(?:UTF-8'')?([^;]+)", cd, re.I)
    if m:
        fname = unquote(m.group(1).strip().strip('"'))
    else:
        m = re.search(r'filename="?([^";]+)"?', cd, re.I)
        if m:
            fname = m.group(1)
    if not ext:
        ext = os.path.splitext(fname)[1].lower() or os.path.splitext(urlparse(url).path)[1].lower()
        if not ext:
            ext = {"application/pdf": ".pdf", "text/csv": ".csv", "image/png": ".png",
                   "image/jpeg": ".jpg", "text/plain": ".txt",
                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
                   }.get(ctype.split(";")[0].strip(), "")
        if not ext:
            die("Could not tell what kind of file the document link is. Download it and give "
                "me the file instead.")
    stem = safe_label(os.path.splitext(fname)[0]) if fname else "document"
    os.makedirs(os.path.join(WORK, "source"), exist_ok=True)
    with _DL_LOCK:  # downloads can run in parallel: pick a free file name and write it atomically
        path, n = os.path.join(WORK, "source", stem + ext), 2
        while path in _DOC_CACHE.values():
            path = os.path.join(WORK, "source", f"{stem}_{n}{ext}")
            n += 1
        with open(path, "wb") as f:
            f.write(r.content)
        _DOC_CACHE[src] = path
    return path


SKIP_LINK_RE = re.compile(r"figma\.com|/drive/folders/", re.I)
DOC_LINK_RE = re.compile(r"docs\.google\.com/document/d/|drive\.google\.com/file/d/", re.I)


def sheet_link_documents(xlsx_path, tabs):
    """Every document link in the chosen tabs of a tracker sheet -> [(label, link)].
    Figma links and Drive folders are skipped. The label is 'row name' or, under a header row of
    hotel / section names, 'header - row name', so each doc is matched to its own page."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    want = {t.strip().lower() for t in tabs}
    found, seen = [], set()
    for ws in wb.worksheets:
        if ws.title.strip().lower() not in want:
            continue
        headers = {}
        for row in ws.iter_rows():
            vals = [("" if c.value is None else str(c.value).strip()) for c in row]
            links = [c.hyperlink.target if c.hyperlink and c.hyperlink.target else "" for c in row]
            if not vals or not any(vals):
                continue
            if not vals[0] and sum(1 for v in vals[1:] if v) >= 2 and not any(links):
                headers = {i: v for i, v in enumerate(vals) if v and i > 0}
                continue
            for i, link in enumerate(links):
                if not link or SKIP_LINK_RE.search(link) or not DOC_LINK_RE.search(link) or link in seen:
                    continue
                seen.add(link)
                name = (vals[0] or vals[i] or "document").splitlines()[0].strip()
                head = headers.get(i, "")
                found.append((f"{head} - {name}" if head else name, link))
    return found


def _sheet_link_items(src, cfg):
    """Expand a Google Sheet / Excel tracker into the documents it links to (config: sheet_link_tabs)."""
    tabs = cfg.get("sheet_link_tabs") or []
    if not tabs:
        return None
    if re.match(r"^https?://", src, re.I):
        if not re.search(r"docs\.google\.com/spreadsheets/", src):
            return None
        path = _download_link(src)
    elif src.lower().endswith((".xlsx", ".xlsm")) and os.path.isfile(src):
        path = src
    else:
        return None
    links = sheet_link_documents(path, tabs)
    if not links:
        die("No document links found in tab(s): " + ", ".join(tabs) + ". Check the tab name.")
    print(f"SHEET_LINKS={len(links)} (Figma links and folders skipped)", flush=True)
    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=max(1, int(cfg.get("concurrency", 4)))) as ex:
        locals_ = list(ex.map(_download_link, [link for _, link in links]))
    items = []
    for (label, link), local in zip(links, locals_):
        dst = os.path.join(WORK, "source", safe_label(label) + os.path.splitext(local)[1])
        if os.path.abspath(local) != os.path.abspath(dst):
            shutil.copyfile(local, dst)
        items.append((safe_label(label), dst))
    return items


_LOCAL_DOCS = {}


def local_documents(cfg):
    """Cached per run: the same documents are asked for several times in one command."""
    key = json.dumps([cfg.get("document"), cfg.get("sheet_link_tabs")], sort_keys=True)
    if key not in _LOCAL_DOCS:
        _LOCAL_DOCS[key] = _local_documents(cfg)
    return list(_LOCAL_DOCS[key])


def _local_documents(cfg):
    """Return [(label, local_path)] for every document QA shared. The label is the file name
    without its extension and is what is matched to a website page."""
    items, used = [], set()

    def add(path):
        label = os.path.splitext(os.path.basename(path))[0]
        base, n = label, 2
        while label in used:
            label = f"{base}_{n}"
            n += 1
        used.add(label)
        items.append((label, path))

    for src in doc_sources(cfg):
        expanded = _sheet_link_items(src, cfg)
        if expanded is not None:
            for label, path in expanded:
                used.add(label)
                items.append((label, path))
        elif re.match(r"^https?://", src, re.I):
            add(_download_link(src))
        elif os.path.isdir(src):
            for f in sorted(os.listdir(src)):
                fp = os.path.join(src, f)
                if f.startswith(".") or f.startswith("~$") or not os.path.isfile(fp):
                    continue
                if not f.lower().endswith(SUPPORTED_DOC_EXT):
                    print(f"SKIPPED_FILE {f} (unsupported file type)")
                    continue
                add(fp)
        elif os.path.exists(src):
            add(src)
        else:
            die(f"Document not found: {src!r}. Check the file name / link you gave.")
    if not items:
        die("No documents found.")
    return items





def source_hash(items):
    h = hashlib.sha256()
    for label, path in items:
        h.update(label.encode())
        with open(path, "rb") as fh:
            h.update(fh.read())
    return h.hexdigest()


def extract_document(items):
    """Return document text. Several files -> '[File: name]' sections. Image / scanned pages
    appear as [[TRANSCRIBE: png]] markers."""
    shutil.rmtree(IMG_DIR, ignore_errors=True)
    counter = [0]
    if len(items) == 1:
        return _extract_one(items[0][1], counter)
    return "\n\n".join(f"[File: {label}]\n{_extract_one(path, counter)}" for label, path in items)


def write_doc_files(text):
    os.makedirs(WORK, exist_ok=True)
    with open(DOC_TXT, "w", encoding="utf-8") as f:
        f.write(text)
    docs_dir = os.path.join(WORK, "docs")  # one text file per shared document
    shutil.rmtree(docs_dir, ignore_errors=True)
    parts = re.split(r"(?m)^\[File: (.+?)\][ \t]*$", text)
    if len(parts) > 1:
        os.makedirs(docs_dir, exist_ok=True)
        for label, body in zip(parts[1::2], parts[2::2]):
            with open(os.path.join(docs_dir, safe_label(label) + ".txt"), "w", encoding="utf-8") as f:
                f.write(body.strip() + "\n")
    chunk_dir = os.path.join(WORK, "doc_chunks")
    os.makedirs(chunk_dir, exist_ok=True)
    for old in os.listdir(chunk_dir):
        os.remove(os.path.join(chunk_dir, old))
    chunks, cur = [], ""
    for para in re.split(r"\n\s*\n", text):
        if len(cur) + len(para) > 4000 and cur:
            chunks.append(cur)
            cur = ""
        cur += para + "\n\n"
    if cur.strip():
        chunks.append(cur)
    for i, c in enumerate(chunks, 1):
        with open(os.path.join(chunk_dir, f"{i:03d}.txt"), "w", encoding="utf-8") as f:
            f.write(c)
    return len(chunks)


def file_hashes():
    d, out = os.path.join(WORK, "docs"), {}
    if os.path.isdir(d):
        for f in os.listdir(d):
            with open(os.path.join(d, f), encoding="utf-8") as fh:
                out[f[:-4]] = sha(fh.read())
    return out


def print_transcribe_request():
    with open(DOC_TXT, encoding="utf-8") as f:
        todo = TRANSCRIBE_RE.findall(f.read())
    print(f"DOC_NEEDS_TRANSCRIPTION={len(todo)}")
    print(f"DOC_TEXT_FILE={DOC_TXT}")
    for p in todo:
        print(f"DOC_IMAGE={p}")


def load_document(cfg, state):
    """Read the document(s). Returns (changed, needs_transcription)."""
    items = local_documents(cfg)
    shash = source_hash(items)
    if state and state.get("doc_source_hash") == shash and os.path.exists(DOC_TXT) \
            and not state.get("awaiting_doc"):
        return False, False  # same files as last time - keep existing text (incl. transcription)
    text = extract_document(items)
    pending = bool(TRANSCRIBE_RE.search(text))
    if not pending and len(text.strip()) < 20:
        die("The document has no readable text.")
    write_doc_files(text)
    return True, pending


# ---------------------------------------------------------------- fetching
EXPAND_JS = """
async () => {
  const sleep = ms => new Promise(r => setTimeout(r, ms));
  const seen = new Set();
  for (let round = 0; round < 3; round++) {
    const targets = [];
    document.querySelectorAll('[aria-expanded="false"], .MuiAccordionSummary-root').forEach(e => {
      if (e.getAttribute('aria-expanded') !== 'true') targets.push(e);
    });
    document.querySelectorAll('img[alt*="more" i], img[alt*="expand" i], img[alt*="plus" i]').forEach(e => {
      targets.push(e.closest('[contenteditable], button, [role=button]') || e);
    });
    document.querySelectorAll('button, [role=button], span, div').forEach(e => {
      if (e.children.length === 0 && /^(read|show|view|see)\\s+more$|^expand( all)?$/i.test(e.textContent.trim()) && !e.closest('a[href]')) targets.push(e);
    });
    let clicked = 0;
    for (const t of targets) {
      if (seen.has(t)) continue;
      seen.add(t);
      try { t.click(); clicked++; } catch (e) {}
    }
    if (!clicked) break;
    await sleep(350);  // let the page open what was clicked (and any nested blocks)
  }
}
"""


_BROWSER_READY = False


def _ensure_browser():
    """Install playwright + Chromium on first use (needed to open collapsed text)."""
    global _BROWSER_READY
    if _BROWSER_READY:
        return
    try:
        import playwright  # noqa: F401
    except ImportError:
        subprocess.call([sys.executable, "-m", "pip", "install", "--quiet", "playwright"])
    subprocess.call([sys.executable, "-m", "playwright", "install", "chromium"],
                    stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    _BROWSER_READY = True


async def _fetch_many_js(urls, timeout, expand, concurrency):
    from playwright.async_api import async_playwright
    results = {}
    sem = asyncio.Semaphore(concurrency)
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        context = await browser.new_context(user_agent=UA)

        async def one(url):
            async with sem:
                page = await context.new_page()
                try:
                    resp = await page.goto(url, wait_until="domcontentloaded", timeout=timeout * 1000)
                    try:  # let the page finish loading, but never wait on endless trackers
                        await page.wait_for_load_state("networkidle", timeout=4000)
                    except Exception:
                        pass
                    if expand:
                        try:
                            await page.evaluate(EXPAND_JS)
                        except Exception:
                            pass
                    results[url] = (resp.status if resp else 0, await page.content(), "")
                except Exception as e:
                    results[url] = (0, "", f"{type(e).__name__}: {e}")
                finally:
                    await page.close()

        await asyncio.gather(*(one(u) for u in urls))
        await browser.close()
    return results


def fetch_many_js(urls, timeout, expand=False, concurrency=4):
    """Open every URL in ONE headless Chromium (several tabs at once). With expand=True, open every
    collapsed block (show more / read more / expand icons, accordions) so hidden text is checked."""
    _ensure_browser()
    return asyncio.run(_fetch_many_js(list(urls), timeout, expand, concurrency))


def fetch_js(url, timeout, expand=False):
    status, html, err = fetch_many_js([url], timeout, expand)[url]
    if err:
        raise RuntimeError(err)
    return status, html


def fetch(url, cfg):
    """Return (status_code, html, error_string)."""
    try:
        if cfg["render_js"] or cfg.get("expand_collapsed"):
            try:
                status, html = fetch_js(url, cfg["timeout_seconds"], expand=cfg.get("expand_collapsed", False))
                return status, html, ""
            except Exception:
                if cfg["render_js"]:
                    raise
                # no browser available: fall back to a plain download (collapsed text stays hidden)
        r = requests.get(url, headers={"User-Agent": UA}, timeout=cfg["timeout_seconds"])
        ctype = r.headers.get("content-type", "")
        if "html" not in ctype and "xml" not in ctype and "text" not in ctype:
            return r.status_code, "", f"non-HTML content-type: {ctype}"
        r.encoding = r.apparent_encoding if r.encoding in (None, "ISO-8859-1") else r.encoding
        return r.status_code, r.text, ""
    except Exception as e:  # network errors etc.
        return 0, "", f"{type(e).__name__}: {e}"


BLOCK_TAGS = {"p", "div", "li", "ul", "ol", "section", "article", "td", "th", "tr", "table", "header",
              "footer", "nav", "main", "aside", "form", "br", "h1", "h2", "h3", "h4", "h5", "h6"}


def _block_of(node):
    for parent in node.parents:
        if parent.name in BLOCK_TAGS:
            return parent
    return None


def _join_inline_sentences(body):
    """Inline pieces of one paragraph (spans) are shown together by the browser. Join a piece to
    the previous one when a sentence ends with no space, so 'luxury.Beautiful' stays visible."""
    nodes = [n for n in body.find_all(string=True) if n.strip()]
    for prev, node in zip(nodes, nodes[1:]):
        if (re.search(r"[.!?]\u200b?$", prev) and re.match(r"[A-Za-z]", node.lstrip("\u200b"))
                and _block_of(prev) is _block_of(node) and _block_of(prev) is not None
                and not prev.endswith((" ", "\n")) and not node.startswith((" ", "\n"))):
            joined = str(prev) + str(node)
            prev.replace_with("")
            node.replace_with(NavigableString(joined))


def extract_page(html, url):
    """Return (text for checking, human page name). Headings are tagged [H1]..[H6]."""
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "noscript", "template", "svg"]):
        t.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    meta = soup.find("meta", attrs={"name": re.compile(r"^description$", re.I)})
    meta_desc = (meta.get("content") or "").strip() if meta else ""
    alts = [i.get("alt").strip() for i in soup.find_all("img") if i.get("alt", "").strip()]
    body = soup.body or soup
    h1 = ""
    for h in body.find_all(re.compile(r"^h[1-6]$")):
        txt = h.get_text(" ", strip=True)
        if not txt:
            continue
        if h.name == "h1" and not h1:
            h1 = txt
        h.clear()
        h.append(NavigableString(f"[{h.name.upper()}] {txt}"))
    _join_inline_sentences(body)
    text = body.get_text("\n", strip=True)
    text = re.sub(r"\n{3,}", "\n\n", text)
    out = [f"URL: {url}", f"TITLE: {title}", f"META_DESCRIPTION: {meta_desc}",
           "--- BODY ---", text]
    if alts:
        out += ["--- IMAGE ALT TEXTS ---"] + alts
    path = urlparse(url).path.strip("/").split("/")[-1].replace("-", " ").replace("_", " ")
    name = h1 or title or path.title() or "Home"
    return "\n".join(out), name


def get_page(url, cfg):
    """Fetch + extract one page. Falls back to a JS-rendered fetch for empty app shells."""
    status, html, err = fetch(url, cfg)
    res = {"status": status, "html": html, "error": err, "text": "", "name": "", "warning": ""}
    if status == 200 and html:
        res["text"], res["name"] = extract_page(html, url)
        if not cfg["render_js"] and not cfg.get("expand_collapsed") and len(res["text"]) < 450 and html.lower().count("<script") >= 3:
            try:
                s2, h2 = fetch_js(url, cfg["timeout_seconds"])
                if s2 == 200 and h2:
                    res["html"] = h2
                    res["text"], res["name"] = extract_page(h2, url)
            except Exception:
                res["warning"] = ("page looks JavaScript-rendered and may be incomplete; "
                                  "set render_js to true and install playwright")
    return res


_ANCHORS = {}   # normalized url -> link texts that point to it (menu labels etc.)
_NAMES = {}     # normalized url -> page title / h1 seen while crawling


def get_pages(urls, cfg):
    """Fetch + extract many pages at once. Returns {url: result}. Plain downloads run in parallel
    threads; pages that need the browser share ONE Chromium with several tabs."""
    urls = list(dict.fromkeys(urls))
    out = {}
    if not urls:
        return out
    workers = max(1, int(cfg.get("concurrency", 4)))
    use_js = cfg["render_js"] or cfg.get("expand_collapsed")
    if use_js:
        try:
            raw = fetch_many_js(urls, cfg["timeout_seconds"], expand=bool(cfg.get("expand_collapsed")),
                                concurrency=workers)
        except Exception:
            if cfg["render_js"]:
                raise
            raw = None  # no browser available: plain download below
        if raw is not None:
            for u in urls:
                status, html, err = raw[u]
                res = {"status": status, "html": html, "error": err, "text": "", "name": "", "warning": ""}
                if status == 200 and html:
                    res["text"], res["name"] = extract_page(html, u)
                out[u] = res
            return out
    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=workers) as ex:
        for u, res in zip(urls, ex.map(lambda x: get_page(x, cfg), urls)):
            out[u] = res
    return out


def extract_links(html, base, cfg=None):
    soup = BeautifulSoup(html, "html.parser")
    if cfg is not None:
        nb = norm_url(base, cfg)
        if nb:
            names = _NAMES.setdefault(nb, set())
            if soup.title and soup.title.get_text(strip=True):
                names.add(soup.title.get_text(" ", strip=True).split("|")[0].strip())
            h1 = soup.find("h1")
            if h1 and h1.get_text(strip=True):
                names.add(h1.get_text(" ", strip=True))
    out = []
    for a in soup.find_all("a", href=True):
        link = urljoin(base, a["href"])
        out.append(link)
        if cfg is not None:
            txt = a.get_text(" ", strip=True)
            nu = norm_url(link, cfg)
            if nu and txt and len(txt) <= 60:
                _ANCHORS.setdefault(nu, set()).add(txt)
    return out


# ------------------------------------------------------------- URL logic
def host_key(netloc):
    n = netloc.lower()
    return n[4:] if n.startswith("www.") else n


def norm_url(url, cfg):
    url, _ = urldefrag(url.strip())
    p = urlparse(url)
    if p.scheme not in ("http", "https"):
        return None
    path = re.sub(r"/index\.(html?|php)$", "/", p.path or "/", flags=re.I)  # /index.html == /
    path = path or "/"
    if path != "/" and path.endswith("/"):
        path = path[:-1]
    query = "" if cfg["strip_query"] else p.query
    out = f"{p.scheme}://{p.netloc.lower()}{path}"
    return out + ("?" + query if query else "")


def allowed_url(url, start, cfg, robots):
    p = urlparse(url)
    if host_key(p.netloc) != host_key(urlparse(start).netloc):
        return False
    if p.path.lower().endswith(SKIP_EXT):
        return False
    for pat in cfg["exclude_patterns"]:
        if re.search(pat, url):
            return False
    if robots and not robots.can_fetch(UA, url):
        return False
    return True


def load_robots(start, cfg):
    if not cfg["respect_robots"]:
        return None
    from urllib.robotparser import RobotFileParser
    p = urlparse(start)
    rp = RobotFileParser()
    try:
        r = requests.get(f"{p.scheme}://{p.netloc}/robots.txt",
                         headers={"User-Agent": UA}, timeout=cfg["timeout_seconds"])
        if r.status_code == 200:
            rp.parse(r.text.splitlines())
            return rp
    except Exception:
        pass
    return None


def sitemap_urls(sm_url, cfg, out, depth=0, seen=None):
    seen = seen if seen is not None else set()
    if depth > 3 or sm_url in seen:
        return
    seen.add(sm_url)
    try:
        r = requests.get(sm_url, headers={"User-Agent": UA}, timeout=cfg["timeout_seconds"])
        if r.status_code != 200:
            return
        root = ET.fromstring(r.content)
    except Exception:
        return
    is_index = root.tag.split("}")[-1] == "sitemapindex"
    for el in root.iter():
        if el.tag.split("}")[-1] == "loc" and el.text:
            loc = el.text.strip()
            if is_index:
                sitemap_urls(loc, cfg, out, depth + 1, seen)
            else:
                out.append(loc)


def discover_urls(cfg, max_depth=None, limit=None):
    """Find EVERY page: follow links from the home page AND merge in the sitemap.

    Using both means pages missing from the sitemap, and pages only reachable by link,
    are all included. Returns (urls, number_dropped_by_max_pages).
    """
    start = norm_url(cfg["url"], cfg)
    if not start:
        die("Invalid website link.")
    cap = limit or cfg["max_pages"]
    depth_cap = cfg["max_depth"] if max_depth is None else max_depth
    if cfg.get("page_urls"):  # QA asked for specific pages only -> no crawling
        listed = [norm_url(urljoin(start + "/", u), cfg) for u in cfg["page_urls"]]
        ordered = []
        for u in listed:
            if u and u not in ordered:
                ordered.append(u)
        return ordered, 0
    robots = load_robots(start, cfg)

    crawled = []
    seen = {start}
    level, depth = [start], 0
    crawl_cfg = dict(cfg, expand_collapsed=False)  # only links are needed here: no browser clicks
    from concurrent.futures import ThreadPoolExecutor
    workers = max(1, int(cfg.get("concurrency", 4)))
    while level and len(crawled) < cap * 2:  # hard stop for runaway sites
        level = [u for u in level if allowed_url(u, start, cfg, robots)][: cap * 2 - len(crawled)]
        crawled.extend(level)
        with ThreadPoolExecutor(max_workers=workers) as ex:
            fetched = list(ex.map(lambda u: fetch(u, crawl_cfg), level))
        nxt = []
        if depth < depth_cap:
            for url, (status, html, _) in zip(level, fetched):
                if status == 200 and html:
                    for link in extract_links(html, url, cfg):
                        nu = norm_url(link, cfg)
                        if nu and nu not in seen:
                            seen.add(nu)
                            nxt.append(nu)
        level, depth = nxt, depth + 1

    from_sitemap = []
    if cfg["use_sitemap"]:
        p = urlparse(start)
        sm_list = [f"{p.scheme}://{p.netloc}/sitemap.xml"]
        if robots and hasattr(robots, "site_maps") and robots.site_maps():
            sm_list = list(robots.site_maps()) + sm_list
        raw = []
        for sm in sm_list:
            sitemap_urls(sm, cfg, raw)
        for u in raw:
            nu = norm_url(u, cfg)
            if nu and allowed_url(nu, start, cfg, robots):
                from_sitemap.append(nu)

    ordered, seen_o = [], set()
    for u in [start] + crawled + from_sitemap:
        if u not in seen_o:
            seen_o.add(u)
            ordered.append(u)
    dropped = max(0, len(ordered) - cap)
    return ordered[:cap], dropped


# ------------------------------------------------- file name -> website page
_STOP = {"global", "page", "pages", "content", "doc", "docx", "final", "copy", "new", "v", "version",
         "option", "options", "the", "of", "and", "for", "website", "web", "text", "copydeck", "draft"}
_HOME_WORDS = {"home", "homepage", "landing", "index", "main"}


def _tokens(s):
    return [t for t in re.findall(r"[a-z0-9]+", s.lower()) if t not in _STOP and not t.isdigit()]


def _score(ltoks, cand):
    ctoks = _tokens(cand)
    if not ctoks or not ltoks:
        return 0.0
    if set(ltoks) == set(ctoks):
        return 1.0
    return difflib.SequenceMatcher(None, " ".join(ltoks), " ".join(ctoks)).ratio()


def match_files(labels, urls, start, existing=None):
    """Match document file names to website pages using the URL path, menu / link text and page
    titles. Returns {label: {"url": url-or-"", "score": float, "how": "auto"}}."""
    out = {}
    for label in labels:
        if existing and existing.get(label, {}).get("url"):
            out[label] = existing[label]
            continue
        lt = _tokens(label)
        if lt and set(lt) <= _HOME_WORDS:
            out[label] = {"url": start, "score": 1.0, "how": "auto"}
            continue
        best = (0.0, "", 99)
        for u in urls:
            if u == start:
                continue
            seg = urlparse(u).path.strip("/").split("/")[-1].replace("-", " ").replace("_", " ")
            depth = len([x for x in urlparse(u).path.split("/") if x])
            for c in ([seg] if seg else []) + list(_ANCHORS.get(u, ())) + list(_NAMES.get(u, ())):
                sc = _score(lt, c)
                if sc > best[0] + 1e-9 or (abs(sc - best[0]) < 1e-9 and sc > 0 and depth < best[2]):
                    best = (sc, u, depth)
        ok = best[0] >= 0.75
        out[label] = {"url": best[1] if ok else "", "score": round(best[0], 2), "how": "auto"}
    return out


# -------------------------------------------------------------- commands
def cmd_setup(args):
    cfg = {}
    if os.path.exists(CONFIG_NAME):
        with open(CONFIG_NAME, encoding="utf-8") as f:
            cfg = json.load(f)
    if not cfg:
        cfg = {"document": "", "url": "", "ignore_terms": []}
    if args.url:
        u = args.url.strip()
        if not re.match(r"^https?://", u, re.I):
            local = re.match(r"^(localhost|\d{1,3}(\.\d{1,3}){3})(:\d+)?(/|$)", u, re.I)
            u = ("http://" if local else "https://") + u
        cfg["url"] = u
    if args.doc:
        docs = [d.strip() for d in args.doc if d.strip()]
        cfg["document"] = docs[0] if len(docs) == 1 else docs
    if args.sheet_links:
        cfg["sheet_link_tabs"] = [t.strip() for t in args.sheet_links.split(",") if t.strip()]
    if args.pages:
        cfg["page_urls"] = [u.strip() for u in args.pages.split(",") if u.strip()]
    if args.all_pages:
        cfg["check_all_pages"] = True
    if args.ignore:
        cfg["ignore_terms"] = [t.strip() for t in args.ignore.split(",") if t.strip()]
    with open(CONFIG_NAME, "w", encoding="utf-8") as f:
        json.dump(cfg, f, indent=2)
    print(f"CONFIG_SAVED url={cfg.get('url')} documents={len(doc_sources(cfg))}")


def new_state(cfg, urls, shash, text):
    return {
        "created": now(), "url": cfg["url"], "document": doc_display(cfg),
        "doc_hash": sha(text), "doc_source_hash": shash, "awaiting_doc": False,
        "multi_doc": False, "file_map": {}, "file_hashes": {},
        "run": 1, "mode": "initial", "batch_counter": 0, "open_batch": None,
        "run_seen": {}, "coverage_run": 0, "coverage": [],
        "pages": {u: {"status": "pending", "hash": None, "checked_run": 0, "name": ""}
                  for u in urls},
        "issues": {},
    }


def print_file_map(state):
    if not state.get("multi_doc"):
        return
    n_ok = sum(1 for m in state["file_map"].values() if m.get("url"))
    print(f"DOCUMENT_FILES={len(state['file_map'])} MATCHED={n_ok}")
    for label, m in state["file_map"].items():
        if m.get("url"):
            print(f"FILE_MATCH {label} -> {m['url']}  (match {int(m.get('score', 0) * 100)}%, {m.get('how', 'auto')})")
        else:
            print(f"FILE_UNMATCHED {label}  (best guess too weak; use: map --file \"{label}\" --url <page>)")


def cmd_prepare(args):
    cfg = load_config({"url": getattr(args, "url", None)})
    if getattr(args, "doc", None):
        cfg["document"] = args.doc
    if os.path.exists(STATE_PATH) and not args.force:
        die("A previous run exists. Use /skillcrawl (re-check) or /skillcrawl --fresh.")
    if not cfg["url"]:
        die("No website link set.")
    changed, pending = load_document(cfg, None)
    items = local_documents(cfg)
    multi = len(items) > 1
    with open(DOC_TXT, encoding="utf-8") as f:
        text = f.read()
    print("Finding the pages on the website (this can take a minute)...", flush=True)
    start = norm_url(cfg["url"], cfg)
    fmap, dropped = {}, 0
    if multi and not cfg["check_all_pages"]:
        urls, _ = discover_urls(cfg, max_depth=cfg["match_depth"], limit=cfg["match_pages"])
        fmap = match_files([l for l, _p in items], urls, start)
        page_list = []
        for m in fmap.values():
            if m["url"] and m["url"] not in page_list:
                page_list.append(m["url"])
    else:
        page_list, dropped = discover_urls(cfg)
        if multi:
            fmap = match_files([l for l, _p in items], page_list, start)
    state = new_state(cfg, page_list, source_hash(items), text)
    state.update({"multi_doc": multi, "file_map": fmap, "file_hashes": file_hashes(),
                  "awaiting_doc": pending})
    save_state(state)
    print(f"DOC_CHARS={len(text)}")
    print(f"PAGES_TO_CHECK={len(page_list)}")
    print(f"BATCHES_EXPECTED={-(-len(page_list) // cfg['batch_size'])}")
    print_file_map(state)
    if dropped:
        print(f"WARNING_PAGES_CAPPED={dropped} more pages exist than max_pages "
              f"({cfg['max_pages']}). Raise max_pages in {CONFIG_NAME} to check them all.")
    if pending:
        print_transcribe_request()


def cmd_map(args):
    """Tell the plugin which website page a document file belongs to (fixes a wrong / missing match)."""
    cfg = load_config()
    state = load_state()
    want = args.file.strip().lower()
    labels = [l for l in state["file_map"] if l.lower() == want or safe_label(l).lower() == want]
    labels = labels or [l for l in state["file_map"] if want in l.lower()]
    if len(labels) != 1:
        die(f"Could not pick one document from {args.file!r}. Documents: {', '.join(state['file_map'])}")
    u = args.url.strip()
    if not re.match(r"^https?://", u, re.I):
        u = urljoin(cfg["url"], u if u.startswith("/") else "/" + u)
    nu = norm_url(u, cfg)
    state["file_map"][labels[0]] = {"url": nu, "score": 1.0, "how": "manual"}
    pg = state["pages"].setdefault(nu, {"status": "pending", "hash": None, "checked_run": 0, "name": ""})
    pg["status"] = "pending"
    save_state(state)
    print(f"MAPPED {labels[0]} -> {nu}")


def pending_urls(state):
    return [u for u, p in state["pages"].items() if p["status"] == "pending"]


def cmd_doc_ready(_):
    state = load_state()
    with open(DOC_TXT, encoding="utf-8") as f:
        text = f.read()
    left = TRANSCRIBE_RE.findall(text)
    if left:
        die(f"{len(left)} document image(s) still not transcribed in {DOC_TXT}: "
            + ", ".join(left[:5]))
    write_doc_files(text)
    state["doc_hash"] = sha(text)
    state["file_hashes"] = file_hashes()
    state["awaiting_doc"] = False
    save_state(state)
    print("DOC_READY")


def cmd_next_batch(_):
    cfg = load_config()
    state = load_state()
    if state.get("awaiting_doc"):
        die("The document images have not been transcribed yet. Run: doc-ready after "
            "transcribing.")

    if state.get("open_batch"):  # resume an unsubmitted batch
        print_batch(state["open_batch"]["dir"], state)
        return

    pend = pending_urls(state)
    if not pend:
        print("NO_MORE_BATCHES")
        return

    urls = pend[: cfg["batch_size"]]
    state["batch_counter"] += 1
    bid = f"run{state['run']:02d}_batch{state['batch_counter']:03d}"
    bdir = os.path.join(WORK, "batches", bid)
    os.makedirs(bdir, exist_ok=True)

    pages = []
    fetched = get_pages(urls, cfg)
    for i, url in enumerate(urls, 1):
        res = fetched[url]
        entry = {"url": url, "name": res["name"], "http_status": res["status"],
                 "error": res["error"], "file": None, "hash": None, "warning": res["warning"],
                 "docs": docs_for(state, url)}
        if res["status"] == 200 and res["text"]:
            fpath = os.path.join(bdir, f"page_{i}.txt")
            with open(fpath, "w", encoding="utf-8") as f:
                f.write(res["text"])
            entry["file"] = fpath
            entry["hash"] = sha(res["text"])
        else:
            entry["error"] = res["error"] or f"HTTP {res['status']}"
        pages.append(entry)

    manifest = {"batch_id": bid, "run": state["run"], "created": now(),
                "ignore_terms": cfg["ignore_terms"], "pages": pages}
    with open(os.path.join(bdir, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)
    state["open_batch"] = {"id": bid, "dir": bdir}
    save_state(state)
    print_batch(bdir, state)


def docs_for(state, url):
    """Document text file(s) this page must be compared with."""
    if not state.get("multi_doc"):
        return [DOC_TXT]
    return [os.path.join(WORK, "docs", safe_label(l) + ".txt")
            for l, m in state["file_map"].items() if m.get("url") == url]


def print_batch(bdir, state):
    with open(os.path.join(bdir, "manifest.json"), encoding="utf-8") as f:
        m = json.load(f)
    print(f"BATCH_DIR={bdir}")
    print(f"BATCH_ID={m['batch_id']}")
    print(f"PAGES_IN_BATCH={len(m['pages'])}")
    print(f"PENDING_TOTAL_INCLUDING_THIS_BATCH={len(pending_urls(state))}")
    print(f"IGNORE_TERMS={json.dumps(m['ignore_terms'])}")
    for p in m["pages"]:
        if p["file"]:
            docs = ",".join(p.get("docs") or []) or "NONE (check spelling / grammar only)"
            print(f"PAGE file={p['file']} url={p['url']} document={docs}")
            if p.get("warning"):
                print(f"PAGE_WARNING url={p['url']} {p['warning']}")
        else:
            print(f"PAGE_ERROR url={p['url']} error={p['error']}")
    print(f"WRITE_ISSUES_TO={os.path.join(bdir, 'issues.json')}")


def make_issue_id(url, itype, found):
    return hashlib.sha1(f"{url}|{itype}|{norm(found).lower()}".encode()).hexdigest()[:8]


def find_section(raw_text, found):
    """Name of the heading / area of the page that contains `found`."""
    if not found:
        return ""
    ranges, off = [], 0
    for ln in raw_text.split("\n"):
        n = norm(ln)
        if not n:
            continue
        ranges.append((off, n))
        off += len(n) + 1
    joined = " ".join(n for _, n in ranges)
    idx = joined.find(norm(found))
    if idx < 0:
        return ""
    sec = ""
    for o, n in ranges:
        if o > idx:
            break
        m = re.match(r"\[H[1-6]\] (.*)", n)
        if m:
            sec = m.group(1)
        elif n.startswith("TITLE:"):
            sec = "Page title (browser tab)"
        elif n.startswith("META_DESCRIPTION:"):
            sec = "Meta description"
        elif n == "--- BODY ---":
            sec = "Top of page"
        elif n == "--- IMAGE ALT TEXTS ---":
            sec = "Image alt text"
    return sec













def find_pos(raw_text, found, section):
    """Position of the issue on the page (lets the report list fixes top-to-bottom)."""
    joined = " ".join(n for n in (norm(l) for l in raw_text.split("\n")) if n)
    if found:
        i = joined.find(norm(found))
        if i >= 0:
            return i
    if section:
        i = joined.find(norm(section))
        if i >= 0:
            return i
    return 10 ** 9


def cmd_submit(args):
    cfg = load_config()
    state = load_state()
    bdir = args.batch_dir
    with open(os.path.join(bdir, "manifest.json"), encoding="utf-8") as f:
        manifest = json.load(f)
    ipath = os.path.join(bdir, "issues.json")
    if not os.path.exists(ipath):
        die(f"{ipath} not found. Write the issues file first (use [] if no issues).")
    try:
        with open(ipath, encoding="utf-8") as f:
            raw_issues = json.load(f)
    except json.JSONDecodeError as e:
        die(f"issues.json is not valid JSON: {e}")
    if isinstance(raw_issues, dict):
        raw_issues = raw_issues.get("issues", [])

    page_raw, page_text, page_name = {}, {}, {}
    for p in manifest["pages"]:
        page_name[p["url"]] = p.get("name", "")
        if p["file"]:
            with open(p["file"], encoding="utf-8") as f:
                page_raw[p["url"]] = f.read()
            page_text[p["url"]] = norm(page_raw[p["url"]])
    run = state["run"]
    ignore = {t.lower() for t in cfg["ignore_terms"]}
    ignore_words = {w for t in ignore for w in re.findall(r"[\w'-]+", t)}
    accepted, dropped, ignored = [], [], []

    # automatic issues for pages that could not be fetched
    for p in manifest["pages"]:
        if not p["file"]:
            accepted.append({"page_url": p["url"], "page_name": p.get("name") or "Broken link: " + (urlparse(p["url"]).path or "/"),
                             "section": "Whole page", "type": "broken_page", "severity": "high",
                             "found_text": f"HTTP {p['http_status']}", "pos": 0,
                             "suggestion": "This page does not load. Fix the page or remove the "
                                           "link that points to it.",
                             "explanation": p["error"]})

    for it in raw_issues:
        url = norm_url(it.get("page_url", ""), cfg) or it.get("page_url", "")
        itype = it.get("type", "other")
        itype = itype if itype in ISSUE_TYPES else "other"
        found = it.get("found_text", "") or ""
        if url not in page_text:
            dropped.append((url, "page not in this batch"))
            continue
        if found and norm(found) not in page_text[url]:
            dropped.append((url, f"found_text not verbatim on page: {found[:60]!r}"))
            continue
        if not found and itype not in ("missing_content", "content_mismatch", "consistency", "other"):
            dropped.append((url, "no found_text"))
            continue
        # brand names / jargon: only a SHORT spelling flag that is (or contains) an ignored term
        words = re.findall(r"[\w'-]+", found.lower())
        if itype == "spelling" and words and all(w in ignore_words for w in words):
            ignored.append((url, found))
            continue
        sev = it.get("severity", "medium")
        sev = sev if sev in SEV_ORDER else "medium"
        section = find_section(page_raw[url], found) or it.get("section", "") or ""
        accepted.append({"page_url": url, "page_name": page_name.get(url, ""), "section": section,
                         "type": itype, "severity": sev, "found_text": found,
                         "pos": find_pos(page_raw[url], found, section),
                         "instruction": it.get("instruction", ""),
                         "suggestion": it.get("suggestion", ""),
                         "source_reference": it.get("source_reference", ""),
                         "explanation": it.get("explanation", "")})

    rows = []
    for it in accepted:
        key = it.get("found_text") or it.get("suggestion", "")
        iid = make_issue_id(it["page_url"], it["type"], key)
        existing = state["issues"].get(iid)
        if existing is None:
            it["first_seen_run"] = run
            label = "new"
        else:
            label = "regressed" if existing["status"] == "fixed" else "still_open"
            it["first_seen_run"] = run if label == "regressed" else existing.get("first_seen_run", run)
        it["status"] = "open"
        it["last_seen_run"] = run
        state["issues"][iid] = it
        state["run_seen"][iid] = label
        rows.append(issue_row(iid, it, label))

    for p in manifest["pages"]:
        pg = state["pages"].setdefault(p["url"], {})
        pg["status"] = "done"
        pg["hash"] = p["hash"]
        pg["name"] = p.get("name", "") or pg.get("name", "")
        pg["checked_run"] = run
    state["open_batch"] = None
    save_state(state)

    report.write_csv(os.path.join(cfg["report_dir"], "batches", manifest["batch_id"] + ".csv"), rows)
    print(f"SUBMITTED {manifest['batch_id']}: {len(rows)} issues recorded "
          f"({sum(1 for r in rows if r['status'] == 'new')} new, "
          f"{sum(1 for r in rows if r['status'] == 'regressed')} reopened)")
    for url, why in dropped:
        print(f"DROPPED_ISSUE {url}: {why}")
    for url, txt in ignored:
        print(f"IGNORED_BY_ignore_terms {url}: {txt!r}")
    print(f"PENDING_PAGES={len(pending_urls(state))}")


def cmd_pages(_):
    state = load_state()
    run = state["run"]
    n = 0
    for u, p in state["pages"].items():
        if p["status"] == "done" or (p["status"] == "pending" and p.get("name")):
            print(f"{p.get('name') or '(no name)'} | {u}")
            n += 1
    print(f"PAGES_LISTED={n}  (doc-vs-website coverage: write {COVERAGE_PATH})")


def record_coverage(state, cov):
    """cov = [{doc_page, url|"" , note}] - one entry per document page / file."""
    run = state["run"]
    known = set(state["pages"])
    missing = 0
    for c in cov:
        c["url"] = (norm_url(c["url"], {"strip_query": True}) or "") if c.get("url") else ""
        if c["url"] and c["url"] not in known:
            c["url"] = ""  # claimed a page we never crawled -> treat as not found
        if c["url"]:
            continue
        name = c.get("doc_page", "").strip()
        if not name:
            continue
        it = {"page_url": MISSING_PAGE_URL, "page_name": name, "section": "",
              "type": "missing_page", "severity": "high", "found_text": "", "pos": 10 ** 9,
              "page_name": re.sub(r"[-_]+", " ", name).strip(),
              "suggestion": f"Create and publish the page \"{name}\". Its content is in the document "
                            f"\"{c.get('note') or name}\". No matching page was found on the website.",
              "source_reference": c.get("note", "") or name,
              "explanation": "Document page with no matching page on the website.",
              "first_seen_run": run, "status": "open", "last_seen_run": run}
        iid = make_issue_id(MISSING_PAGE_URL, "missing_page", name)
        existing = state["issues"].get(iid)
        label = "new" if existing is None else (
            "regressed" if existing["status"] == "fixed" else "still_open")
        if existing is not None and label == "still_open":
            it["first_seen_run"] = existing.get("first_seen_run", run)
        state["issues"][iid] = it
        state["run_seen"][iid] = label
        missing += 1
    state["coverage"] = cov
    state["coverage_run"] = run
    return missing


def cmd_coverage_submit(_):
    state = load_state()
    if not os.path.exists(COVERAGE_PATH):
        die(f"{COVERAGE_PATH} not found. Write it (use [] if the document has no page list).")
    with open(COVERAGE_PATH, encoding="utf-8") as f:
        cov = json.load(f)
    missing = record_coverage(state, cov)
    save_state(state)
    print(f"COVERAGE_SAVED document_pages={len(cov)} missing_on_website={missing}")


# ---------------------------------------------------------------- reports






def cmd_finalize(_):
    cfg = load_config()
    state = load_state()
    run = state["run"]
    if state.get("open_batch"):
        die("There is an unsubmitted batch. Submit it first.")
    pend = pending_urls(state)
    if pend:
        die(f"{len(pend)} pages still pending. Keep running next-batch.")
    if state.get("multi_doc"):  # coverage comes straight from the file -> page matching
        record_coverage(state, [{"doc_page": l, "url": m.get("url", ""), "note": l}
                                for l, m in state["file_map"].items()])
    elif state.get("coverage_run") != run:
        print("WARNING_COVERAGE_NOT_DONE: run `pages`, write coverage.json, then "
              "`coverage-submit` so missing document pages are reported.")

    fixed_ids = []
    for iid, it in state["issues"].items():
        if it["status"] != "open" or iid in state["run_seen"]:
            continue
        if it["type"] == "missing_page":
            if state.get("coverage_run") != run:
                continue  # coverage not re-evaluated this run -> leave as is
        else:
            pg = state["pages"].get(it["page_url"], {})
            if pg.get("checked_run") != run:
                continue
        it["status"] = "fixed"
        it["fixed_run"] = run
        fixed_ids.append(iid)
    save_state(state)

    rdir = cfg["report_dir"]
    keep = report.read_dev_columns(rdir)
    cov_name = ({c["url"]: re.sub(r"[-_]+", " ", c["doc_page"]).strip()
                 for c in state.get("coverage", [])
                 if c.get("url") and c.get("doc_page")} if state.get("coverage_run") == run else {})
    rows = []
    for iid, it in state["issues"].items():
        if it["status"] != "open":
            continue
        label = state["run_seen"].get(iid, "still_open")
        r = issue_row(iid, it, STATUS_LABEL[label], keep)
        r["raw_status"] = label
        # show pages under the names used in the document (e.g. "Homepage"), not the HTML heading
        if r["page_url"] in cov_name:
            r["page_name"] = cov_name[r["page_url"]]
        rows.append(r)
    counts = {k: sum(1 for r in rows if r["raw_status"] == k)
              for k in ("new", "still_open", "regressed")}
    fixed_rows = [issue_row(i, state["issues"][i], "FIXED") for i in fixed_ids]
    for r in fixed_rows:
        r["page_name"] = cov_name.get(r["page_url"], r["page_name"])

    # group by page, in the order the pages were listed / checked; number the items 1..N
    order = {u: i for i, u in enumerate(state["pages"])}
    by_page = {}
    for r in sorted(rows, key=lambda r: order.get(r["page_url"], 10 ** 6)):
        g = by_page.setdefault(r["page_url"], {"name": r["page_name"] or r["page_url"],
                                               "url": r["page_url"], "items": []})
        g["items"].append(r)
    groups, n = [], 0
    for g in by_page.values():
        ordered = report.order_rows(g["items"])
        for r in ordered:
            n += 1
            r["no"] = n
        g["fix"] = [r for r in ordered if report.action_of(r) not in ("ADD", "CREATE PAGE")]
        g["add"] = [r for r in ordered if report.action_of(r) in ("ADD", "CREATE PAGE")]
        groups.append(g)
    rows = [r for g in groups for r in g["fix"] + g["add"]]

    open_by_page = {g["url"]: len(g["fix"]) + len(g["add"]) for g in groups}
    pages_table = [(cov_name.get(u) or p.get("name") or u, u, open_by_page.get(u, 0))
                   for u, p in state["pages"].items() if p["status"] == "done"]
    covered = {c.get("url") for c in state.get("coverage", []) if c.get("url")}
    extra_pages = ([(p.get("name") or u, u) for u, p in state["pages"].items()
                    if p["status"] == "done" and u not in covered]
                   if state.get("coverage_run") == run and state.get("coverage")
                   and not state.get("multi_doc") else [])
    pages_checked = sum(1 for p in state["pages"].values() if p.get("checked_run") == run)
    stats = {"total": len(rows), "high": sum(1 for r in rows if r["severity"] == "high"),
             "new": counts["new"], "still_open": counts["still_open"],
             "regressed": counts["regressed"], "fixed": len(fixed_ids), "pages": pages_checked}
    title = f"Website fix report - run {run} ({'first check' if run == 1 else 're-check'})"
    meta = [f"Website: {state['url']}", f"Document(s): {doc_display(cfg)}", f"Generated: {now()}"]
    all_clear = not rows
    pages_summary = [{"name": g["name"], "fix": len(g["fix"]), "add": len(g["add"]),
                      "high": sum(1 for r in g["fix"] + g["add"] if r["severity"] == "high"),
                      "total": len(g["fix"]) + len(g["add"])} for g in groups]

    html_doc = report.build_html(title, meta, stats, groups, fixed_rows, pages_table, extra_pages, all_clear)
    md_doc = report.build_md(title, meta, stats, groups, fixed_rows, all_clear)
    hist = os.path.join(rdir, "history", f"run{run:02d}")
    for d in (rdir, hist):
        os.makedirs(d, exist_ok=True)
        with open(os.path.join(d, "FIX_REPORT.html"), "w", encoding="utf-8") as f:
            f.write(html_doc)
        with open(os.path.join(d, "FIX_REPORT.md"), "w", encoding="utf-8") as f:
            f.write(md_doc)
        report.write_csv(os.path.join(d, "FIX_REPORT.csv"), rows)
        report.write_xlsx(os.path.join(d, "FIX_REPORT.xlsx"), meta, pages_summary, rows)

    line = (f"Run {run}: ALL CLEAR ({len(fixed_ids)} fixed)" if all_clear else
            f"Run {run}: {len(rows)} things to fix ({counts['new']} new, "
            f"{counts['still_open']} still open, {counts['regressed']} reopened), "
            f"{len(fixed_ids)} fixed. Report: {os.path.join(rdir, 'FIX_REPORT.html')}")
    if cfg["notify_webhook"]:
        try:
            requests.post(cfg["notify_webhook"], json={"text": line}, timeout=15)
        except Exception as ex:
            print(f"(webhook failed: {ex})")
    print("ALL_CLEAR" if all_clear else f"OPEN_ISSUES={len(rows)}")
    print(f"PAGES_READ={pages_checked} PAGES_ON_SITE={len(state['pages'])}")
    print(f"NEW={counts['new']} STILL_OPEN={counts['still_open']} REOPENED={counts['regressed']} "
          f"FIXED={len(fixed_ids)}")
    print(line)
    for ext in ("html", "xlsx", "csv", "md"):
        print(f"REPORT_{ext.upper()}={os.path.join(rdir, 'FIX_REPORT.' + ext)}")


def start_recheck(cfg, state):
    """Prepare a re-check run. Returns False if the document needs transcription first."""
    changed, pending = load_document(cfg, state)
    items = local_documents(cfg)
    if changed:
        state["doc_source_hash"] = source_hash(items)
        state["force_all"] = True
    if pending:
        state["awaiting_doc"] = True
        save_state(state)
        print_transcribe_request()
        return False
    with open(DOC_TXT, encoding="utf-8") as f:
        text = f.read()
    multi = state.get("multi_doc", False)
    forced = state.pop("force_all", False)
    doc_changed = forced if multi else (forced or sha(text) != state.get("doc_hash"))
    state["doc_hash"] = sha(text)

    state["run"] += 1
    run = state["run"]
    state["mode"] = "recheck"
    state["run_seen"] = {}
    state["batch_counter"] = 0

    force_urls, dropped = set(), 0
    if multi:
        start = norm_url(cfg["url"], cfg)
        labels = [l for l, _p in items]
        fmap = {l: m for l, m in state.get("file_map", {}).items() if l in labels}
        todo = [l for l in labels if not fmap.get(l, {}).get("url")]
        if todo:  # new or still-unmatched files: look for their pages again
            urls0, _ = discover_urls(cfg, max_depth=cfg["match_depth"], limit=cfg["match_pages"])
            fmap.update(match_files(todo, urls0, start))
        for l in labels:
            fmap.setdefault(l, {"url": "", "score": 0, "how": "auto"})
        state["file_map"] = fmap
        if cfg["check_all_pages"]:
            urls, dropped = discover_urls(cfg)
        else:
            urls = []
            for m in fmap.values():
                if m["url"] and m["url"] not in urls:
                    urls.append(m["url"])
        cur = file_hashes()
        old = state.get("file_hashes", {})
        changed_labels = {safe_label(l) for l in labels if old.get(safe_label(l)) != cur.get(safe_label(l))}
        force_urls = {m["url"] for l, m in fmap.items() if m["url"] and safe_label(l) in changed_labels}
        state["file_hashes"] = cur
    else:
        urls, dropped = discover_urls(cfg)
    discovered = set(urls)
    for u in urls:
        state["pages"].setdefault(u, {"status": "pending", "hash": None, "checked_run": 0,
                                      "name": ""})

    open_pages = {it["page_url"] for it in state["issues"].values() if it["status"] == "open"}
    n_pending = n_skipped = n_gone = 0
    to_fetch = [u for u, pg in state["pages"].items() if pg["status"] != "pending"
                and not (multi and not cfg["check_all_pages"] and u not in discovered)]
    fetched = get_pages(to_fetch, cfg)
    for url, pg in state["pages"].items():
        if pg["status"] == "pending":
            n_pending += 1
            continue
        if multi and not cfg["check_all_pages"] and url not in discovered:
            pg["status"], pg["checked_run"] = "gone", run  # no document is mapped to it any more
            n_gone += 1
            continue
        res = fetched[url]
        if res["status"] in (404, 410) and url not in discovered:
            # removed from the site AND nothing links to it any more -> closed
            pg["status"], pg["checked_run"] = "gone", run
            n_gone += 1
            continue
        new_hash = sha(res["text"]) if res["status"] == 200 and res["text"] else None
        changed_page = new_hash != pg.get("hash")
        if doc_changed or changed_page or url in open_pages or url in force_urls \
                or pg["status"] != "done":
            pg["status"] = "pending"
            n_pending += 1
        else:
            pg["checked_run"] = run  # unchanged and clean -> counts as checked
            n_skipped += 1
    save_state(state)
    print(f"RECHECK_RUN={run}")
    print(f"DOCUMENT_CHANGED={doc_changed or bool(force_urls)}")
    print(f"PAGES_TO_CHECK={n_pending} UNCHANGED_SKIPPED={n_skipped} GONE={n_gone}")
    print_file_map(state)
    if dropped:
        print(f"WARNING_PAGES_CAPPED={dropped} more pages exist than max_pages.")
    return True


def cmd_auto(args):
    """One entry point for QA: decides first run / resume / re-check by itself."""
    if not os.path.exists(CONFIG_NAME):
        print("NEEDS_SETUP: ask the user for the website link and the document path(s), "
              "then run: setup --url <link> --doc <path> [<path> ...]")
        return
    cfg = load_config()
    if not cfg["url"] or not doc_sources(cfg):
        print("NEEDS_SETUP: website link or document missing in config.")
        return
    state = None
    if os.path.exists(STATE_PATH) and not args.fresh:
        state = load_state()
        multi_now = len(local_documents(cfg)) > 1
        if state.get("url") != cfg["url"] or multi_now != state.get("multi_doc", False):
            print("NOTE: the website or the number of documents changed - starting a new full check.")
            state = None
    if state is None:
        print("MODE=FIRST_RUN")
        cmd_prepare(argparse.Namespace(url=None, doc=None, force=True))
        return
    if state.get("awaiting_doc"):
        print("MODE=NEEDS_DOCUMENT_TRANSCRIPTION")
        print_transcribe_request()
        return
    if state.get("open_batch") or pending_urls(state):
        print("MODE=RESUME")
        print(f"PAGES_PENDING={len(pending_urls(state))}")
        print_file_map(state)
        return
    print("MODE=RECHECK")
    start_recheck(cfg, state)


def cmd_status(_):
    state = load_state()
    by = {}
    for p in state["pages"].values():
        by[p["status"]] = by.get(p["status"], 0) + 1
    open_n = sum(1 for i in state["issues"].values() if i["status"] == "open")
    fixed_n = sum(1 for i in state["issues"].values() if i["status"] == "fixed")
    print(f"run={state['run']} mode={state['mode']} pages={by} open_issues={open_n} "
          f"fixed_issues={fixed_n} open_batch={'yes' if state.get('open_batch') else 'no'}")


def main():
    ap = argparse.ArgumentParser(description="SkillCrawl engine")
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("setup")
    s.add_argument("--url")
    s.add_argument("--doc", nargs="+", help="one or more files / folders / links")
    s.add_argument("--all-pages", action="store_true",
                   help="with several files: also check every other page of the site")
    s.add_argument("--pages", help="comma-separated page URLs / paths to check (instead of all)")
    s.add_argument("--sheet-links", help="comma-separated tab name(s) of a tracker sheet: open every "
                   "document link in those tabs (Figma links skipped) and check the site against them")
    s.add_argument("--ignore", help="comma-separated brand names / terms to never flag")
    s.set_defaults(fn=cmd_setup)
    m = sub.add_parser("map")
    m.add_argument("--file", required=True, help="document file name (as listed by auto)")
    m.add_argument("--url", required=True, help="website page URL or path, e.g. /about-us")
    m.set_defaults(fn=cmd_map)
    a = sub.add_parser("auto")
    a.add_argument("--fresh", action="store_true")
    a.set_defaults(fn=cmd_auto)
    sub.add_parser("next-batch").set_defaults(fn=cmd_next_batch)
    s = sub.add_parser("submit")
    s.add_argument("batch_dir")
    s.set_defaults(fn=cmd_submit)
    sub.add_parser("doc-ready").set_defaults(fn=cmd_doc_ready)
    sub.add_parser("pages").set_defaults(fn=cmd_pages)
    sub.add_parser("coverage-submit").set_defaults(fn=cmd_coverage_submit)
    sub.add_parser("finalize").set_defaults(fn=cmd_finalize)
    sub.add_parser("status").set_defaults(fn=cmd_status)
    args = ap.parse_args()
    args.fn(args)


if __name__ == "__main__":
    main()
