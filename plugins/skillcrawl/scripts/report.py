"""Developer-facing fix reports for SkillCrawl: HTML, Excel, CSV and Markdown.

Design goals: a developer who has never seen the plugin can open the report and know, for every
item, WHICH page, WHERE on the page, WHAT is wrong, and EXACTLY what to change it to.
"""
import csv
import difflib
import html as htmllib
import os
import re

TYPE_LABEL = {"spelling": "Spelling mistake", "grammar": "Grammar mistake",
              "punctuation": "Punctuation mistake",
              "content_mismatch": "Text does not match the document",
              "missing_content": "Missing from the page",
              "missing_page": "Page missing from the website",
              "consistency": "Inconsistent text", "broken_page": "Page not loading",
              "other": "Other"}
ACTION_OF_TYPE = {"missing_content": "ADD", "missing_page": "CREATE PAGE",
                  "broken_page": "FIX PAGE", "consistency": "REVIEW", "other": "REVIEW"}
ACTION_LABEL = {"REPLACE": "Replace text", "ADD": "Add content", "CREATE PAGE": "Create page",
                "FIX PAGE": "Fix page", "REVIEW": "Review & decide"}
DEFAULT_INSTRUCTION = {
    "REPLACE": "Find the text in the red box on this page and replace it with the text in the "
               "green box.",
    "ADD": "Add the text in the green box to this page, in the section named above.",
    "CREATE PAGE": "Create and publish this page on the website. Its content is in the document named below.",
    "FIX PAGE": "This page does not load. Fix it, or remove the links that point to it.",
    "REVIEW": "Check this item and decide what is correct, then update the page or the "
              "document.",
}
SEV_ORDER = {"high": 0, "medium": 1, "low": 2}
SEV_LABEL = {"high": "HIGH", "medium": "MEDIUM", "low": "LOW"}

FIX_COLUMNS = ["No.", "Page", "Section", "Action", "Priority", "Currently on site", "Change to",
               "What to do", "Why", "Document reference", "Status", "Page link", "Issue ID",
               "Assignee", "Dev status", "Dev notes"]
DEV_COLUMNS = ("Assignee", "Dev status", "Dev notes")


# ------------------------------------------------------------------ helpers
def action_of(r):
    return ACTION_OF_TYPE.get(r.get("type"), "REPLACE")


def instruction_of(r):
    ins = (r.get("instruction") or "").strip()
    if ins:
        return ins
    act, found, sug = action_of(r), r.get("found_text", ""), r.get("suggestion", "")
    if act == "REPLACE" and found and sug and len(found) <= 70 and len(sug) <= 70:
        return f"Change \u201c{found}\u201d to \u201c{sug}\u201d."
    return DEFAULT_INSTRUCTION[act]


def issue_row(iid, it, label, keep=None):
    dev = (keep or {}).get(iid, {})
    return {"issue_id": iid, "status": label, "severity": it.get("severity", "medium"),
            "type": it["type"], "page_name": it.get("page_name", ""), "page_url": it["page_url"],
            "section": it.get("section", ""), "found_text": it.get("found_text", ""),
            "suggestion": it.get("suggestion", ""), "instruction": it.get("instruction", ""),
            "source_reference": it.get("source_reference", ""),
            "explanation": it.get("explanation", ""), "first_seen_run": it.get("first_seen_run", ""),
            "pos": it.get("pos", 10 ** 9),
            "assignee": dev.get(DEV_COLUMNS[0], ""), "dev_status": dev.get(DEV_COLUMNS[1], ""),
            "dev_notes": dev.get(DEV_COLUMNS[2], "")}


def _toks(s):
    return re.findall(r"\s+|[\w’'-]+|[^\w\s]", s)


def diff_html(old, new):
    """Word-level diff: returns (old_html, new_html) with <del>/<ins> around the changed words.
    Returns None when the two texts are too different for a diff to help."""
    a, b = _toks(old), _toks(new)
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    if sm.ratio() < 0.35:
        return None
    e = htmllib.escape
    o, n = [], []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        A, B = "".join(a[i1:i2]), "".join(b[j1:j2])
        if tag == "equal":
            o.append(e(A))
            n.append(e(B))
            continue
        o.append(f"<del>{e(A)}</del>" if A.strip() else e(A))
        n.append(f"<ins>{e(B)}</ins>" if B.strip() else e(B))
    return "".join(o), "".join(n)


def order_rows(rows):
    """Developer order: text fixes top-to-bottom as they appear on the page, then additions."""
    return sorted(rows, key=lambda r: (0 if action_of(r) in ("REPLACE", "REVIEW", "FIX PAGE") else 1,
                                       r.get("pos", 10 ** 9), SEV_ORDER.get(r["severity"], 1)))


# --------------------------------------------------------------------- CSV
def _fix_row(r):
    return [r.get("no", ""), r.get("page_name", ""), r.get("section", "") or "-",
            ACTION_LABEL[action_of(r)], SEV_LABEL.get(r["severity"], r["severity"]),
            r.get("found_text", ""), r.get("suggestion", ""), instruction_of(r),
            r.get("explanation", ""), r.get("source_reference", ""), r.get("status", ""),
            r.get("page_url", ""), r.get("issue_id", ""), r.get("assignee", ""),
            r.get("dev_status", ""), r.get("dev_notes", "")]


def write_csv(path, rows):
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig: opens cleanly in Excel
        w = csv.writer(f)
        w.writerow(FIX_COLUMNS)
        for r in rows:
            w.writerow(_fix_row(r))


def read_dev_columns(folder):
    """Keep what developers typed (assignee / status / notes) across re-runs (xlsx or csv)."""
    keep = {}
    xl = os.path.join(folder, "FIX_REPORT.xlsx")
    if os.path.exists(xl):
        try:
            import openpyxl
            ws = openpyxl.load_workbook(xl, read_only=True)["Fixes"]
            head = [c for c in next(ws.iter_rows(values_only=True))]
            for row in list(ws.iter_rows(values_only=True))[1:]:
                d = dict(zip(head, row))
                if d.get("Issue ID"):
                    keep[str(d["Issue ID"])] = {k: str(d.get(k) or "") for k in DEV_COLUMNS}
        except Exception:
            pass
    cs = os.path.join(folder, "FIX_REPORT.csv")
    if os.path.exists(cs):
        try:
            with open(cs, newline="", encoding="utf-8-sig") as f:
                for d in csv.DictReader(f):
                    iid = d.get("Issue ID")
                    got = {k: d.get(k, "") for k in DEV_COLUMNS}
                    if iid and any(got.values()):
                        base = keep.get(iid, {})
                        keep[iid] = {k: base.get(k) or got[k] for k in DEV_COLUMNS}
        except Exception:
            pass
    return keep


# -------------------------------------------------------------------- XLSX
def write_xlsx(path, meta_lines, pages_summary, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "SkillCrawl fix report"
    ws["A1"].font = Font(size=16, bold=True)
    r = 2
    for m in meta_lines:
        ws.cell(row=r, column=1, value=m)
        r += 1
    r += 1
    heads = ["Page", "Fix text", "Add content", "High priority", "Total"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2A44")
    for p in pages_summary:
        r += 1
        for c, v in enumerate([p["name"], p["fix"], p["add"], p["high"], p["total"]], 1):
            ws.cell(row=r, column=c, value=v)
    ws.column_dimensions["A"].width = 44
    for col in "BCDE":
        ws.column_dimensions[col].width = 14

    fx = wb.create_sheet("Fixes")
    fx.append(FIX_COLUMNS)
    for c in range(1, len(FIX_COLUMNS) + 1):
        cell = fx.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2A44")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    fills = {"HIGH": "FDE2E2", "MEDIUM": "FFF0CC", "LOW": "E4ECFF"}
    thin = Side(style="thin", color="D9DDE5")
    for r_ in rows:
        fx.append(_fix_row(r_))
        row = fx.max_row
        pri = fx.cell(row=row, column=5)
        pri.fill = PatternFill("solid", fgColor=fills.get(str(pri.value), "FFFFFF"))
        pri.font = Font(bold=True)
        fx.cell(row=row, column=6).fill = PatternFill("solid", fgColor="FFF1F1")
        fx.cell(row=row, column=7).fill = PatternFill("solid", fgColor="EDFAEF")
        for c in range(1, len(FIX_COLUMNS) + 1):
            cell = fx.cell(row=row, column=c)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    widths = [6, 22, 26, 15, 10, 48, 48, 40, 34, 30, 12, 34, 10, 14, 14, 30]
    for i, w in enumerate(widths, 1):
        fx.column_dimensions[get_column_letter(i)].width = w
    fx.freeze_panes = "C2"
    fx.auto_filter.ref = fx.dimensions
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    return True


# -------------------------------------------------------------------- HTML
CSS = """
:root{color-scheme:light}*{box-sizing:border-box}
body{font-family:system-ui,-apple-system,Segoe UI,Roboto,Arial,sans-serif;margin:0;background:#f3f4f7;
color:#1b2233;line-height:1.5}
main{max-width:1040px;margin:0 auto;padding:28px 16px 80px}
h1{font-size:26px;margin:0 0 6px}h2{font-size:21px;margin:0}h3{font-size:14px;letter-spacing:.06em;
text-transform:uppercase;color:#4a5570;margin:22px 0 8px}
.meta{color:#5b6475;font-size:14px}.meta div{margin:1px 0}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(120px,1fr));gap:10px;margin:20px 0}
.stat{background:#fff;border:1px solid #dfe3ea;border-radius:10px;padding:12px 14px}
.stat b{display:block;font-size:26px;line-height:1.1}.stat span{font-size:12px;color:#5b6475}
.stat.hi b{color:#b3261e}.stat.ok b{color:#1b7a3b}
.howto{background:#fff;border:1px solid #dfe3ea;border-left:5px solid #2f5bd8;border-radius:10px;
padding:14px 18px;margin:16px 0;font-size:14px}.howto ol{margin:6px 0 0 18px;padding:0}
table.toc{width:100%;border-collapse:collapse;background:#fff;border:1px solid #dfe3ea;
border-radius:10px;overflow:hidden;font-size:14px}
.toc th,.toc td{padding:9px 12px;text-align:left;border-top:1px solid #e8ebf0}
.toc th{background:#1f2a44;color:#fff;font-size:12px;text-transform:uppercase;letter-spacing:.04em;
border:0}.toc td.n{text-align:center;width:90px}
.page{background:#fff;border:1px solid #dfe3ea;border-radius:12px;margin:26px 0;padding:18px 18px 8px}
.page-head{display:flex;flex-wrap:wrap;gap:8px 14px;align-items:baseline;justify-content:space-between;
border-bottom:2px solid #1f2a44;padding-bottom:10px}
.page-head a{font-size:13px;color:#2f5bd8;word-break:break-all}
.fix{border:1px solid #e2e6ee;border-left:6px solid #9aa4b8;border-radius:10px;margin:12px 0;padding:12px 14px}
.fix.sev-high{border-left-color:#d93025}.fix.sev-medium{border-left-color:#f0a30a}
.fix.sev-low{border-left-color:#2f5bd8}.fix.isdone{opacity:.45}
.top{display:flex;flex-wrap:wrap;align-items:center;gap:8px}
.no{font-weight:800;font-size:15px;background:#1f2a44;color:#fff;border-radius:6px;padding:1px 8px}
.chip{display:inline-block;font-size:11.5px;font-weight:700;letter-spacing:.03em;border-radius:99px;
padding:2px 10px}
.a-replace{background:#e6efff;color:#1f4fb0}.a-add{background:#e1f5e5;color:#1b6b34}
.a-create-page{background:#e1f5e5;color:#1b6b34}.a-fix-page{background:#fde2e2;color:#a31515}
.a-review{background:#f1e6ff;color:#5b2aa6}
.p-high{background:#fde2e2;color:#a31515}.p-medium{background:#fff0cc;color:#8a5a00}
.p-low{background:#e4ecff;color:#1f4fb0}
.s-new{background:#e1f5e5;color:#1b6b34}.s-still_open{background:#eee;color:#444}
.s-regressed{background:#fde2e2;color:#a31515}
.done{margin-left:auto;font-size:13px;color:#4a5570;cursor:pointer;user-select:none}
.where{margin:8px 0 2px;font-size:14px}.where b{color:#4a5570}
.what{font-weight:650;margin:2px 0 10px;font-size:15px}
.cols{display:grid;grid-template-columns:1fr 1fr;gap:10px}
.cols.one{grid-template-columns:1fr}
@media(max-width:760px){.cols{grid-template-columns:1fr}}
.box{border-radius:8px;padding:10px 12px;border:1px solid}.box .lbl{font-size:11px;font-weight:800;
letter-spacing:.07em;margin-bottom:4px}.box .txt{white-space:pre-wrap;word-break:break-word;font-size:14.5px}
.box.now{background:#fff3f3;border-color:#f1b8b4}.box.now .lbl{color:#a31515}
.box.new{background:#effaf1;border-color:#b3e0bd}.box.new .lbl{color:#1b6b34}
del{background:#ffc9c4;text-decoration:line-through;border-radius:3px}
ins{background:#a9e8b6;text-decoration:none;font-weight:700;border-radius:3px}
.small{font-size:12.5px;color:#5b6475;margin-top:8px}.small b{color:#3a445c}
.empty{color:#8a93a6;font-style:italic}
details{margin:14px 0}summary{cursor:pointer;font-weight:700}
ul.plain{margin:6px 0 0 18px;padding:0;font-size:14px}
@media print{body{background:#fff}.page,.fix{break-inside:avoid}.done{display:none}}
"""

JS = """
try{document.querySelectorAll('input[data-id]').forEach(function(c){
 var k='skillcrawl_done_'+c.dataset.id,a=c.closest('article');
 try{c.checked=localStorage.getItem(k)==='1'}catch(e){}
 a.classList.toggle('isdone',c.checked);
 c.addEventListener('change',function(){try{localStorage.setItem(k,c.checked?'1':'0')}catch(e){}
  a.classList.toggle('isdone',c.checked)})})}catch(e){}
"""


def _box(label, cls, inner_html):
    return f"<div class='box {cls}'><div class='lbl'>{label}</div><div class='txt'>{inner_html}</div></div>"


def _item_html(r):
    e = htmllib.escape
    act = action_of(r)
    sev = r["severity"]
    sec = r.get("section") or ""
    where = f"<b>Where on the page:</b> {e(sec)}" if sec else "<b>Where on the page:</b> whole page"
    found, sug = r.get("found_text", ""), r.get("suggestion", "")
    if act == "REPLACE" and found and sug:
        d = diff_html(found, sug)
        old_h, new_h = d if d else (e(found), e(sug))
        cols = ("<div class='cols'>" + _box("CURRENTLY ON THE SITE", "now", old_h)
                + _box("CHANGE IT TO", "new", new_h) + "</div>")
    elif act == "ADD":
        cols = "<div class='cols one'>" + _box("ADD THIS TEXT", "new", e(sug) or "<span class='empty'>see document</span>") + "</div>"
    elif found:
        cols = ("<div class='cols'>" + _box("CURRENTLY ON THE SITE", "now", e(found))
                + _box("SUGGESTED / ACTION", "new", e(sug)) + "</div>")
    else:
        cols = "<div class='cols one'>" + _box("ACTION", "new", e(sug)) + "</div>"
    small = []
    if r.get("explanation"):
        small.append(f"<b>Why:</b> {e(r['explanation'])}")
    if r.get("source_reference"):
        small.append(f"<b>Source in document:</b> {e(r['source_reference'])}")
    small.append(f"<b>Problem type:</b> {e(TYPE_LABEL.get(r['type'], r['type']))}")
    small.append(f"<b>ID:</b> {e(r['issue_id'])}")
    st = r.get("raw_status", "new")
    return (f"<article class='fix sev-{e(sev)}' id='fix-{r.get('no', '')}'>"
            f"<div class='top'><span class='no'>#{r.get('no', '')}</span>"
            f"<span class='chip a-{act.lower().replace(' ', '-')}'>{e(ACTION_LABEL[act].upper())}</span>"
            f"<span class='chip p-{e(sev)}'>{e(SEV_LABEL.get(sev, sev))} PRIORITY</span>"
            f"<span class='chip s-{e(st)}'>{e(r.get('status', ''))}</span>"
            f"<label class='done'><input type='checkbox' data-id='{e(r['issue_id'])}'> Done</label></div>"
            f"<div class='where'>{where}</div><div class='what'>{e(instruction_of(r))}</div>"
            f"{cols}<div class='small'>{' &nbsp;|&nbsp; '.join(small)}</div></article>")


def build_html(title, meta_lines, stats, groups, fixed, pages_table, extra_pages, all_clear):
    e = htmllib.escape
    h = ["<!doctype html><html lang='en'><head><meta charset='utf-8'>"
         "<meta name='viewport' content='width=device-width,initial-scale=1'>"
         f"<title>{e(title)}</title><style>{CSS}</style></head><body><main>",
         f"<h1>{e(title)}</h1><div class='meta'>"
         + "".join(f"<div>{e(m)}</div>" for m in meta_lines) + "</div>"]
    h.append("<div class='stats'>"
             f"<div class='stat hi'><b>{stats['total']}</b><span>things to fix</span></div>"
             f"<div class='stat hi'><b>{stats['high']}</b><span>high priority</span></div>"
             f"<div class='stat'><b>{stats['new']}</b><span>new</span></div>"
             f"<div class='stat'><b>{stats['still_open']}</b><span>still open</span></div>"
             f"<div class='stat'><b>{stats['regressed']}</b><span>reopened</span></div>"
             f"<div class='stat ok'><b>{stats['fixed']}</b><span>fixed since last run</span></div>"
             f"<div class='stat'><b>{stats['pages']}</b><span>pages checked</span></div></div>")
    if all_clear:
        h.append("<div class='howto' style='border-left-color:#1b7a3b'><b>ALL CLEAR.</b> Every "
                 "checked page matches the document. Nothing for developers to fix.</div>")
    else:
        h.append("<div class='howto'><b>How to use this report</b><ol>"
                 "<li>Find your page below and click <b>Open page</b>.</li>"
                 "<li>For each numbered item, go to the <b>section</b> named on the page.</li>"
                 "<li>Change the text in the <span style='background:#ffc9c4'>red box</span> to "
                 "the text in the <span style='background:#a9e8b6'>green box</span> (changed words "
                 "are highlighted). For <b>Add content</b>, paste the green text into the page.</li>"
                 "<li>Tick <b>Done</b> as you go. Tell QA when finished - they will re-run the check."
                 "</li></ol></div>")
        h.append("<h3>Pages and what needs doing</h3><table class='toc'><tr><th>Page</th>"
                 "<th class='n'>Fix text</th><th class='n'>Add content</th><th class='n'>High</th>"
                 "<th class='n'>Total</th></tr>")
        for i, g in enumerate(groups, 1):
            high = sum(1 for r in g["fix"] + g["add"] if r["severity"] == "high")
            h.append(f"<tr><td><a href='#page-{i}'>{e(g['name'])}</a></td>"
                     f"<td class='n'>{len(g['fix'])}</td><td class='n'>{len(g['add'])}</td>"
                     f"<td class='n'>{high}</td><td class='n'><b>{len(g['fix']) + len(g['add'])}</b></td></tr>")
        h.append("</table>")
    for i, g in enumerate(groups, 1):
        link = (f"<a href='{e(g['url'])}' target='_blank' rel='noopener'>Open page &#8599; {e(g['url'])}</a>"
                if g["url"].startswith("http") else f"<span class='small'>{e(g['url'])}</span>")
        h.append(f"<section class='page' id='page-{i}'><div class='page-head'><h2>{e(g['name'])}</h2>"
                 f"{link}</div>")
        if g["fix"]:
            h.append(f"<h3>Fix text already on the page ({len(g['fix'])})</h3>")
            h.extend(_item_html(r) for r in g["fix"])
        if g["add"]:
            h.append(f"<h3>Add what is missing ({len(g['add'])})</h3>")
            h.extend(_item_html(r) for r in g["add"])
        h.append("</section>")
    if fixed:
        h.append(f"<details open><summary>Fixed since last run ({len(fixed)})</summary><ul class='plain'>")
        for r in fixed:
            h.append(f"<li><b>{e(r['page_name'])}</b> - {e(r['section'] or 'whole page')}: "
                     f"<del>{e((r['found_text'] or r['suggestion'])[:160])}</del></li>")
        h.append("</ul></details>")
    h.append(f"<details><summary>Pages checked ({len(pages_table)}) - confirm nothing was skipped"
             "</summary><table class='toc'><tr><th>Page</th><th>Link</th><th class='n'>Open items</th></tr>")
    for name, url, n in pages_table:
        h.append(f"<tr><td>{e(name)}</td><td><a href='{e(url)}' target='_blank' rel='noopener'>{e(url)}</a></td>"
                 f"<td class='n'>{n if n else 'OK'}</td></tr>")
    h.append("</table></details>")
    if extra_pages:
        h.append("<details><summary>Website pages not mentioned in the document (information only)"
                 "</summary><ul class='plain'>")
        for name, url in extra_pages:
            h.append(f"<li>{e(name)} - {e(url)}</li>")
        h.append("</ul></details>")
    h.append(f"</main><script>{JS}</script></body></html>")
    return "\n".join(h)


# ---------------------------------------------------------------- Markdown
def build_md(title, meta_lines, stats, groups, fixed, all_clear):
    L = [f"# {title}", ""] + [f"- {m}" for m in meta_lines] + [""]
    L.append(f"**{stats['total']} things to fix** ({stats['high']} high priority) - "
             f"{stats['new']} new, {stats['still_open']} still open, {stats['regressed']} reopened. "
             f"Fixed since last run: {stats['fixed']}.")
    L.append("")
    if all_clear:
        L.append("**ALL CLEAR - nothing to fix.**")
    for g in groups:
        L += ["---", f"## {g['name']}", f"Page: {g['url']}", ""]
        for title_, items in (("Fix text already on the page", g["fix"]), ("Add what is missing", g["add"])):
            if not items:
                continue
            L += [f"### {title_} ({len(items)})", ""]
            for r in items:
                act = ACTION_LABEL[action_of(r)]
                L.append(f"#### #{r.get('no', '')} - {act} - {SEV_LABEL.get(r['severity'])} priority - {r.get('status', '')}")
                L.append(f"- **Where:** {r.get('section') or 'whole page'}")
                L.append(f"- **What to do:** {instruction_of(r)}")
                if r.get("found_text"):
                    L.append(f"- **Currently on the site:** {r['found_text']}")
                L.append(f"- **{'Add this text' if action_of(r) == 'ADD' else 'Change it to'}:** {r.get('suggestion', '')}")
                if r.get("explanation"):
                    L.append(f"- **Why:** {r['explanation']}")
                if r.get("source_reference"):
                    L.append(f"- **Source in document:** {r['source_reference']}")
                L.append(f"- ID: {r['issue_id']}")
                L.append("")
    if fixed:
        L += ["---", f"## Fixed since last run ({len(fixed)})"]
        L += [f"- {r['page_name']} / {r['section'] or 'whole page'}: {(r['found_text'] or r['suggestion'])[:160]}"
              for r in fixed]
    return "\n".join(L)
